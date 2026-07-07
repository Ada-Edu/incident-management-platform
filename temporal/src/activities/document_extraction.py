from __future__ import annotations
"""Activities for the Document Field Extraction feature.

Pipeline (see docs/plans/document-field-extraction-plan.md):
  fetch_document  -> extract_fields (THE MODEL CALL) -> persist_extraction
Deterministic guard-rail outcomes (unsupported/too-large/protected/corrupt/no-text)
are returned as {"ok": False, "reason": ...} so Temporal does NOT retry them; the
workflow routes them to persist_failure. Transient IO/model errors raise and are
retried by the activity's retry policy.

Text extraction stays in customer-invisible plumbing here; the spec is behaviour-only.
"""
import io
import json
import logging
import math
import re
from typing import Any

from temporalio import activity
from temporalio.exceptions import ApplicationError

from .. import supabase_rest
from ..config import settings
from ..model_client import converse_text

logger = logging.getLogger(__name__)

MAX_PAGES = 20
MAX_SOURCE_CHARS = 100_000            # hard cap fed to the model
_CHARS_PER_PAGE = 1800                # docx page estimate (no fixed pages in the format)
_STORAGE_BUCKET = "incident-docs"

_SYSTEM_PROMPT = (
    "You extract structured fields from a document. "
    "Return ONLY a single JSON object, no prose, matching exactly:\n"
    '{"parties":[{"name":string,"role":string|null}],'
    '"key_dates":[{"label":string,"date":string}],'
    '"key_terms":[{"label":string,"value":string}]}\n'
    "Rules: list EVERY named party with its role where the document states one, else role=null. "
    "Key terms are monetary amounts, obligations, service levels, durations/notice periods, and "
    "reference numbers EXPLICITLY present. If a section has no items, return an empty array for it. "
    "Never invent values not supported by the text."
)


class ExtractionRejected(Exception):
    """A deterministic, non-retryable guard-rail outcome (maps to a spec edge case)."""

    def __init__(self, reason: str, page_count: int | None = None):
        super().__init__(reason)
        self.reason = reason
        self.page_count = page_count


def _empty_result() -> dict[str, list]:
    return {"parties": [], "key_dates": [], "key_terms": []}


def _extract_text(file_name: str, data: bytes) -> tuple[str, int]:
    """Return (text, page_count) or raise ExtractionRejected for guard-rail cases."""
    ext = (file_name.rsplit(".", 1)[-1] if "." in file_name else "").lower()

    if ext == "pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                # An empty-password decrypt failing => genuinely protected.
                try:
                    if reader.decrypt("") == 0:
                        raise ExtractionRejected("password_protected")
                except ExtractionRejected:
                    raise
                except Exception:
                    raise ExtractionRejected("password_protected")
            page_count = len(reader.pages)
            if page_count > MAX_PAGES:
                raise ExtractionRejected("too_large", page_count)
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        except ExtractionRejected:
            raise
        except (PdfReadError, Exception) as exc:  # noqa: BLE001
            raise ExtractionRejected("corrupt") from exc

    elif ext == "docx":
        import docx  # python-docx
        try:
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(c.text for c in row.cells)
            text = "\n".join(parts)
        except Exception as exc:  # noqa: BLE001 -- unopenable/locked docx
            raise ExtractionRejected("corrupt") from exc
        page_count = max(1, math.ceil(len(text) / _CHARS_PER_PAGE))
        if page_count > MAX_PAGES:
            raise ExtractionRejected("too_large", page_count)

    else:
        raise ExtractionRejected("unsupported_type")

    if not text.strip():
        # PDF/DOCX opened but no readable text (e.g. scanned image).
        raise ExtractionRejected("no_text", page_count)

    return text[:MAX_SOURCE_CHARS], page_count


# ---------------------------------------------------------------------------
# Activity 1 — fetch + guard rails
# ---------------------------------------------------------------------------
@activity.defn
def fetch_document(extraction_id: str) -> dict[str, Any]:
    """Download the document and extract text. Transient IO errors raise (retried);
    deterministic rejections return ok=False (not retried)."""
    row = supabase_rest.get_extraction(extraction_id)
    if row is None:
        # Not transient: nothing to work on.
        raise ApplicationError(f"extraction {extraction_id} not found", non_retryable=True)

    data = supabase_rest.download_object(_STORAGE_BUCKET, row["storage_path"])  # may raise -> retried
    try:
        text, page_count = _extract_text(row.get("file_name") or row["storage_path"], data)
    except ExtractionRejected as rej:
        logger.info("extraction rejected extraction_id=%s reason=%s", extraction_id, rej.reason)
        return {"ok": False, "reason": rej.reason, "page_count": rej.page_count}

    logger.info("extraction fetched extraction_id=%s pages=%s chars=%s", extraction_id, page_count, len(text))
    return {"ok": True, "text": text, "page_count": page_count}


# ---------------------------------------------------------------------------
# Activity 2 — THE MODEL CALL
# ---------------------------------------------------------------------------
@activity.defn
def extract_fields(text: str) -> dict[str, Any]:
    """Call Bedrock to extract parties/key_dates/key_terms. Boto exceptions
    propagate (retried per policy). Malformed JSON after one reparse is a
    non-retryable model_error."""
    raw = converse_text(_SYSTEM_PROMPT, text)
    cleaned = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
    try:
        parsed = json.loads(cleaned)
    except json.JSONDecodeError:
        # One salvage attempt: grab the outermost JSON object.
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if not match:
            raise ApplicationError("model returned non-JSON", type="model_error", non_retryable=True)
        try:
            parsed = json.loads(match.group(0))
        except json.JSONDecodeError as exc:
            raise ApplicationError("model returned malformed JSON", type="model_error", non_retryable=True) from exc

    result = _empty_result()
    for group in result:
        items = parsed.get(group, [])
        if isinstance(items, list):
            result[group] = items
    return {"result": result, "model_id": settings.bedrock_model_id}


# ---------------------------------------------------------------------------
# Activity 3 — persist success / failure (idempotent upsert on unique id)
# ---------------------------------------------------------------------------
@activity.defn
def persist_extraction(extraction_id: str, result: dict[str, Any], page_count: int | None, model_id: str) -> None:
    supabase_rest.update_extraction(extraction_id, {
        "status": "succeeded",
        "result": result,
        "page_count": page_count,
        "model_id": model_id,
        "failure_reason": None,
    })
    logger.info("extraction persisted extraction_id=%s status=succeeded", extraction_id)


@activity.defn
def persist_failure(extraction_id: str, reason: str, page_count: int | None) -> None:
    supabase_rest.update_extraction(extraction_id, {
        "status": "failed",
        "failure_reason": reason,
        "page_count": page_count,
        "result": _empty_result(),
    })
    logger.info("extraction persisted extraction_id=%s status=failed reason=%s", extraction_id, reason)


@activity.defn
def mark_running(extraction_id: str) -> None:
    supabase_rest.update_extraction(extraction_id, {"status": "running"})
